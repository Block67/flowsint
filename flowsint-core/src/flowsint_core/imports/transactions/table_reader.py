"""
Read tabular transaction files (CSV / XLSX) into headers + rows.
"""

import csv
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List

ALLOWED_TABLE_EXTENSIONS = [".csv", ".xlsx"]

# Rows scanned to find the header row: statements often start with a few
# title lines (bank name, account, period) before the real header.
HEADER_SCAN_ROWS = 20
# A row is the header when it fills at least this share of the widest row.
HEADER_MIN_FILL_RATIO = 0.6


@dataclass
class TableData:
    headers: List[str]
    rows: List[Dict[str, Any]]
    # 1-based line (CSV) / row (Excel) number of each row in the original file
    line_numbers: List[int]


def read_table(file_bytes: bytes, filename: str) -> TableData:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        raw_rows = _read_csv_rows(file_bytes)
    elif ext == ".xlsx":
        raw_rows = _read_xlsx_rows(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file format: {ext}. Supported: {', '.join(ALLOWED_TABLE_EXTENSIONS)}"
        )
    return _to_table(raw_rows)


def _decode(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("latin-1")


def _read_csv_rows(file_bytes: bytes) -> List[List[Any]]:
    text = _decode(file_bytes)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _read_xlsx_rows(file_bytes: bytes) -> List[List[Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Could not read Excel file: {e}")
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _to_table(raw_rows: List[List[Any]]) -> TableData:
    numbered = [
        (line, r)
        for line, r in enumerate(raw_rows, start=1)
        if any(not _is_empty(v) for v in r)
    ]
    if not numbered:
        raise ValueError("The file contains no data.")

    # Header = first row filled enough compared to the widest of the first
    # rows (title lines above it have only one or two cells).
    filled = [sum(not _is_empty(v) for v in r) for _, r in numbered[:HEADER_SCAN_ROWS]]
    threshold = max(filled) * HEADER_MIN_FILL_RATIO
    header_idx = next(i for i, count in enumerate(filled) if count >= threshold)

    headers: List[str] = []
    for i, value in enumerate(numbered[header_idx][1]):
        name = str(value).strip() if not _is_empty(value) else f"column_{i + 1}"
        base, n = name, 2
        while name in headers:
            name = f"{base}_{n}"
            n += 1
        headers.append(name)

    rows: List[Dict[str, Any]] = []
    line_numbers: List[int] = []
    for line, raw in numbered[header_idx + 1 :]:
        values = [v.strip() if isinstance(v, str) else v for v in raw]
        rows.append(
            {h: values[i] if i < len(values) else None for i, h in enumerate(headers)}
        )
        line_numbers.append(line)
    return TableData(headers=headers, rows=rows, line_numbers=line_numbers)
